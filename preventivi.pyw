import sys
import os
import ctypes
import threading
import json
import ssl
import time
import subprocess
import shutil
from datetime import datetime
import tkinter as tk
from tkinter import filedialog

# --- LIBRERIA FONDAMENTALE PER I THREAD DI WORD ---
try:
    import pythoncom
except ImportError:
    pass

# --- 1. GESTIONE ERRORI AVVIO ---
def show_error_msg(title, msg):
    try:
        ctypes.windll.user32.MessageBoxW(0, msg, title, 0x10)
    except:
        print(msg)

# --- 2. IMPORTAZIONE LIBRERIE ---
try:
    import ttkbootstrap as ttk
    from ttkbootstrap.constants import *
    from ttkbootstrap.dialogs import Messagebox
except ImportError:
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "ttkbootstrap"])
        import ttkbootstrap as ttk
        from ttkbootstrap.constants import *
        from ttkbootstrap.dialogs import Messagebox
    except Exception as e:
        show_error_msg("Errore", f"Manca ttkbootstrap: {e}")
        sys.exit()

try:
    from docxtpl import DocxTemplate
    import win32com.client
    from geopy.geocoders import Nominatim
    from geopy.distance import geodesic
    import certifi
    LIBRERIE_OK = True
except ImportError as e:
    LIBRERIE_OK = False
    ERRORE_LIB = str(e)

# --- 3. PERCORSI ---
if getattr(sys, 'frozen', False):
    app_path = os.path.dirname(sys.executable)
elif __file__:
    app_path = os.path.dirname(os.path.abspath(__file__))

path_modello = os.path.join(app_path, "modello_offerta.docx")
path_logo    = os.path.join(app_path, "logo studio serra.png")
path_config  = os.path.join(app_path, "impostazioni.json")
path_icona   = os.path.join(app_path, "icona.ico")
path_bilanci = os.path.join(app_path, "bilanci.json")

# --- 4. CONFIGURAZIONE ---
DEFAULT_CONFIG = {
    "nome_studio": "STUDIO SERRA",
    "indirizzo_studio": "Via Sylva 5, Bergamo",
    "prezzo_fascia_1": 100.0, "prezzo_fascia_2": 80.0,
    "prezzo_fascia_3": 70.0,  "prezzo_fascia_4": 60.0,
    "costo_cancelleria": 100.0, "costo_anagrafica": 50.0,
    "costo_fiscali": 250.0, "costo_cu": 250.0, "costo_km": 1.5,
    "ultima_cartella": "",
    "numero_preventivo": 1,   # M2 – numero progressivo
}

config = DEFAULT_CONFIG.copy()
if os.path.exists(path_config):
    try:
        with open(path_config, 'r') as f:
            config.update(json.load(f))
    except:
        pass

def salva_configurazione():
    try:
        config["nome_studio"]       = ent_cfg_nome.get()
        config["indirizzo_studio"]  = ent_cfg_addr.get()
        config["prezzo_fascia_1"]   = float(ent_cfg_p1.get())
        config["prezzo_fascia_2"]   = float(ent_cfg_p2.get())
        config["prezzo_fascia_3"]   = float(ent_cfg_p3.get())
        config["prezzo_fascia_4"]   = float(ent_cfg_p4.get())
        config["costo_cancelleria"] = float(ent_cfg_canc.get())
        config["costo_anagrafica"]  = float(ent_cfg_anag.get())
        config["costo_fiscali"]     = float(ent_cfg_fisc.get())
        config["costo_cu"]          = float(ent_cfg_cu.get())
        config["costo_km"]          = float(ent_cfg_km.get())   # M5
        with open(path_config, 'w') as f:
            json.dump(config, f, indent=4)
        Messagebox.show_info("Salvato", "Impostazioni aggiornate!")
    except ValueError:
        Messagebox.show_error("Errore", "Inserisci numeri validi.")

# --- 5. LOGICA ---
def mostra_caricamento(titolo):
    top = ttk.Toplevel()
    top.title("")
    w, h = 400, 150
    x = app.winfo_x() + (app.winfo_width()  // 2) - (w // 2)
    y = app.winfo_y() + (app.winfo_height() // 2) - (h // 2)
    top.geometry(f"{w}x{h}+{int(x)}+{int(y)}")
    top.grab_set()
    ttk.Label(top, text=titolo, font=("Helvetica", 10, "bold"), wraplength=380).pack(pady=20)
    pb = ttk.Progressbar(top, mode='indeterminate', length=300)
    pb.pack(fill=X, padx=20)
    pb.start(15)
    return top

def converti_pdf_diretto(input_docx, output_pdf):
    word = None
    try:
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False
        abs_input  = os.path.abspath(input_docx)
        abs_output = os.path.abspath(output_pdf)
        doc = word.Documents.Open(abs_input)
        doc.SaveAs(abs_output, FileFormat=17)
        doc.Close()
        return True
    except Exception as e:
        print(f"Errore Word: {e}")
        return False
    finally:
        try:
            if word: word.Quit()
        except:
            pass

# ─────────────────────────────────────────────────────────
# Calcola e aggiorna il riepilogo importi
# ─────────────────────────────────────────────────────────
def _leggi_float(entry_widget):
    try:
        val = entry_widget.get().replace('€', '').replace(' ', '').strip()
        if ',' in val and '.' in val:
            val = val.replace('.', '').replace(',', '.')
        else:
            val = val.replace(',', '.')
        return float(val) if val else 0.0
    except:
        return 0.0

def aggiorna_riepilogo(*args):
    """Ricalcola i totali in tempo reale e aggiorna i label del pannello riepilogo."""
    try:
        unita  = _leggi_float(entry_unita)
        box    = _leggi_float(entry_box)
        prz    = _leggi_float(entry_prezzo_unita)
        sc     = _leggi_float(entry_sconto)
        spesa  = _leggi_float(entry_spesa)
        asc    = _leggi_float(entry_ascensori)
        canc_t = _leggi_float(entry_cancelli)
        km     = _leggi_float(entry_km)
        c1     = _leggi_float(entry_canc)
        c2     = _leggi_float(entry_anag)
        c3     = _leggi_float(entry_adem)
        c4     = _leggi_float(entry_cu)

        base    = (unita + (box / 5)) * prz
        fattore = var_risc.get() + asc + canc_t + var_port.get() + var_pisc.get()
        extra_b = 0.001 * spesa * fattore
        extra   = extra_b / 2 if spesa > 200000 else extra_b
        lordo   = base + extra
        netto   = lordo - (lordo * (sc / 100))
        trasf   = km * 4 * 2 * config["costo_km"]
        imp     = netto + trasf + c1 + c2 + c3 + c4
        inar    = imp * 0.04
        iva_imp = (imp + inar) * 0.22
        totale  = imp + inar + iva_imp

        lbl_riepilogo_imponibile.config(text=f"€ {imp:,.2f}")
        lbl_riepilogo_cassa.config(text=f"€ {inar:,.2f}")
        lbl_riepilogo_iva.config(text=f"€ {iva_imp:,.2f}")
        lbl_riepilogo_totale.config(text=f"€ {totale:,.2f}")
    except:
        pass

def calcola_tariffa_automatica(event=None):
    try:
        val = entry_unita.get()
        if not val:
            return
        n = int(val)
        if n == 0:
            if entry_prezzo_unita.get() != "0.00":
                entry_prezzo_unita.delete(0, END)
                entry_prezzo_unita.insert(0, "0.00")
            try:
                lbl_tariffa.config(text="")
            except:
                pass
            aggiorna_riepilogo()
            return
        if n < 10:      p, t = config["prezzo_fascia_1"], "Fascia 1 (<10)"
        elif n <= 20:   p, t = config["prezzo_fascia_2"], "Fascia 2 (10-20)"
        elif n <= 30:   p, t = config["prezzo_fascia_3"], "Fascia 3 (21-30)"
        else:           p, t = config["prezzo_fascia_4"], "Fascia 4 (>30)"

        nuovo_p = f"{p:.2f}"
        if entry_prezzo_unita.get() != nuovo_p:
            entry_prezzo_unita.delete(0, END)
            entry_prezzo_unita.insert(0, nuovo_p)
        try:
            if lbl_tariffa.cget("text") != t:
                lbl_tariffa.config(text=t)
        except:
            pass
    except:
        pass
    aggiorna_riepilogo()

def avvia_mappe():
    if not LIBRERIE_OK:
        Messagebox.show_error("Errore", f"Librerie mancanti: {ERRORE_LIB}")
        return
    loading  = mostra_caricamento("Calcolo distanza...")
    addr_s   = config["indirizzo_studio"]
    addr_c   = f"{entry_indirizzo.get()}, {entry_citta.get()}"

    def task():
        err = None; val = 0.0
        try:
            ctx   = ssl.create_default_context(cafile=certifi.where())
            geo   = Nominatim(user_agent="app_serra_v41", ssl_context=ctx)
            loc_s = geo.geocode(addr_s, timeout=10)
            loc_c = geo.geocode(addr_c, timeout=10)
            if loc_s and loc_c:
                dist = geodesic(
                    (loc_s.latitude, loc_s.longitude),
                    (loc_c.latitude, loc_c.longitude)
                ).km
                val = dist * 1.3
            else:
                err = "Indirizzo non trovato."
        except Exception as e:
            err = str(e)
        app.after(0, lambda: fine_mappe(loading, err, val))

    threading.Thread(target=task, daemon=True).start()

def fine_mappe(win, err, val):
    win.destroy()
    if err:
        Messagebox.show_error("Errore", err)
    else:
        entry_km.delete(0, END)
        entry_km.insert(0, f"{val:.1f}")
        Messagebox.show_info("Fatto", f"Distanza: {val:.1f} km")
    aggiorna_riepilogo()

def nuovo_preventivo():
    for e in (entry_nome, entry_indirizzo):
        e.delete(0, END)
    entry_citta.delete(0, END);  entry_citta.insert(0, "Bergamo")
    entry_spesa.delete(0, END);  entry_spesa.insert(0, "0")
    entry_ascensori.delete(0, END); entry_ascensori.insert(0, "0")
    entry_cancelli.delete(0, END);  entry_cancelli.insert(0, "0")
    entry_unita.delete(0, END);  entry_unita.insert(0, "")
    entry_box.delete(0, END);   entry_box.insert(0, "0")
    entry_km.delete(0, END);    entry_km.insert(0, "0")
    entry_prezzo_unita.delete(0, END); entry_prezzo_unita.insert(0, "0")
    entry_sconto.delete(0, END); entry_sconto.insert(0, "0")
    entry_note.delete("1.0", END)
    var_risc.set(0); var_port.set(0); var_pisc.set(0)
    lbl_tariffa.config(text="")
    for w, key in [(entry_canc, "costo_cancelleria"), (entry_anag, "costo_anagrafica"),
                   (entry_adem, "costo_fiscali"),      (entry_cu,  "costo_cu")]:
        w.delete(0, END); w.insert(0, str(config[key]))
    aggiorna_riepilogo()

# ─────────────────────────────────────────────
# M1 – Anteprima preventivo prima del salvataggio
# ─────────────────────────────────────────────
def mostra_anteprima(d, netto, trasf, imp, inar, iva_imp, totale, on_confirm):
    """Mostra riepilogo definitivo e chiede conferma prima di procedere al salvataggio."""
    dlg = ttk.Toplevel()
    dlg.title("📋 Anteprima Preventivo")
    dlg.resizable(False, False)
    w, h = 470, 500
    x = app.winfo_x() + (app.winfo_width()  // 2) - (w // 2)
    y = app.winfo_y() + (app.winfo_height() // 2) - (h // 2)
    dlg.geometry(f"{w}x{h}+{int(x)}+{int(y)}")
    dlg.grab_set()

    ttk.Label(dlg, text="📋  Riepilogo Preventivo",
              font=("Helvetica", 13, "bold"), bootstyle="primary").pack(pady=(15, 5))

    # Dati condominio
    fr_info = ttk.LabelFrame(dlg, text=" Condominio ", padding=8)
    fr_info.pack(fill=X, padx=15, pady=5)
    ttk.Label(fr_info, text=f"Nome:      {d['nome']}", font=("Arial", 10)).pack(anchor=W)
    ttk.Label(fr_info, text=f"Indirizzo: {d['ind']}, {d['cit']}", font=("Arial", 10)).pack(anchor=W)

    # Importi
    fr_imp = ttk.LabelFrame(dlg, text=" Importi ", padding=8)
    fr_imp.pack(fill=X, padx=15, pady=5)
    righe = [
        ("Importo base (netto sconto):", f"€ {netto:,.2f}",   ""),
        ("Trasferte:",                   f"€ {trasf:,.2f}",   ""),
        ("Totale Imponibile:",           f"€ {imp:,.2f}",     "info"),
        ("Cassa Previdenza 4%:",         f"€ {inar:,.2f}",    "warning"),
        ("IVA 22%:",                     f"€ {iva_imp:,.2f}", "secondary"),
        ("TOTALE FINALE:",               f"€ {totale:,.2f}",  "success"),
    ]
    for lbl_txt, val_txt, stile in righe:
        fr_r = ttk.Frame(fr_imp); fr_r.pack(fill=X, pady=2)
        bold = stile in ("success", "info")
        ttk.Label(fr_r, text=lbl_txt,
                  font=("Arial", 10, "bold" if bold else "normal")).pack(side=LEFT)
        lbl_kw = {"bootstyle": stile} if stile else {}
        ttk.Label(fr_r, text=val_txt,
                  font=("Arial", 10, "bold"), **lbl_kw).pack(side=RIGHT)

    # Note (se presenti)
    if d.get("note"):
        fr_note = ttk.LabelFrame(dlg, text=" Note ", padding=8)
        fr_note.pack(fill=X, padx=15, pady=5)
        ttk.Label(fr_note, text=d["note"], wraplength=410,
                  font=("Arial", 9), justify=LEFT).pack(anchor=W)

    # Numero progressivo
    num = config.get("numero_preventivo", 1)
    ttk.Label(dlg, text=f"N° preventivo: {num:03d}",
              font=("Arial", 9, "italic"), bootstyle="secondary").pack(pady=(5, 0))

    # Pulsanti
    fr_btn = ttk.Frame(dlg); fr_btn.pack(fill=X, padx=15, pady=12)

    def conferma():
        dlg.destroy()
        on_confirm()

    ttk.Button(fr_btn, text="✅  Conferma e Salva", command=conferma,
               bootstyle="success", width=22).pack(side=LEFT, padx=5, expand=True, fill=X)
    ttk.Button(fr_btn, text="✏️  Modifica", command=dlg.destroy,
               bootstyle="secondary-outline", width=14).pack(side=RIGHT, padx=5)

# ─────────────────────────────────────────────
# Generazione PDF
# ─────────────────────────────────────────────
def avvia_pdf():
    if not os.path.exists(path_modello):
        Messagebox.show_error("Errore Grave",
            f"Manca il file 'modello_offerta.docx' nella cartella:\n{app_path}")
        return
    if not LIBRERIE_OK:
        Messagebox.show_error("Errore", f"Librerie mancanti (pywin32?): {ERRORE_LIB}")
        return
    try:
        def f(x): return float(x.replace(',', '.') if x else 0)
        d = {
            "nome": entry_nome.get(),
            "ind":  entry_indirizzo.get(),
            "cit":  entry_citta.get(),
            "spesa": f(entry_spesa.get()),
            "asc":   f(entry_ascensori.get()),
            "canc":  f(entry_cancelli.get()),
            "unita": f(entry_unita.get()),
            "box":   f(entry_box.get()),
            "km":    f(entry_km.get()),
            "prz":   f(entry_prezzo_unita.get()),
            "sc":    f(entry_sconto.get()),
            "c1":    f(entry_canc.get()),
            "c2":    f(entry_anag.get()),
            "c3":    f(entry_adem.get()),
            "c4":    f(entry_cu.get()),
            "note":  entry_note.get("1.0", END).strip(),
        }
    except:
        Messagebox.show_error("Errore", "Controlla i numeri inseriti.")
        return

    base    = (d["unita"] + (d["box"] / 5)) * d["prz"]
    fattore = var_risc.get() + d["asc"] + d["canc"] + var_port.get() + var_pisc.get()
    extra_b = 0.001 * d["spesa"] * fattore
    extra   = extra_b / 2 if d["spesa"] > 200000 else extra_b
    lordo   = base + extra
    netto   = lordo - (lordo * (d["sc"] / 100))
    trasf   = d["km"] * 4 * 2 * config["costo_km"]
    imp     = netto + trasf + d["c1"] + d["c2"] + d["c3"] + d["c4"]
    inar    = imp * 0.04
    iva_imp = (imp + inar) * 0.22
    totale  = imp + inar + iva_imp

    # M1 – mostra anteprima prima di procedere
    mostra_anteprima(d, netto, trasf, imp, inar, iva_imp, totale,
                     lambda: _salva_preventivo(d, netto, trasf, imp, inar, iva_imp, totale))

def _salva_preventivo(d, netto, trasf, imp, inar, iva_imp, totale):
    """Esegue il salvataggio effettivo dopo la conferma nell'anteprima."""
    # M2 – numero progressivo nel nome file suggerito
    ultima = config.get("ultima_cartella", "")
    if not ultima or not os.path.isdir(ultima):
        ultima = app_path

    num = config.get("numero_preventivo", 1)
    file_iniziale = f"Preventivo_{num:03d}_{d['nome']}.docx"
    percorso_salvataggio = filedialog.asksaveasfilename(
        title="Dove salvare il preventivo?",
        initialdir=ultima,
        initialfile=file_iniziale,
        defaultextension=".docx",
        filetypes=[("Word Document", "*.docx")]
    )
    if not percorso_salvataggio:
        return

    config["ultima_cartella"] = os.path.dirname(percorso_salvataggio)
    try:
        with open(path_config, 'w') as fh:
            json.dump(config, fh, indent=4)
    except:
        pass

    percorso_pdf = percorso_salvataggio.replace(".docx", ".pdf")
    loading = mostra_caricamento("Generazione Word & PDF...\n(Motore Diretto)")

    def task():
        try:
            pythoncom.CoInitialize()
        except:
            pass
        err     = None
        pdf_ok  = False
        word_ok = False
        try:
            ctx = {
                'nome_condominio':          d["nome"],
                'indirizzo':                d["ind"],
                'citta':                    d["cit"],
                'data_oggi':                datetime.now().strftime("%d/%m/%Y"),
                'importo_base':             f"€ {netto:,.2f}",
                'compenso_annuale':         f"€ {netto + trasf:,.2f}",
                'spese_cancelleria':        f"€ {d['c1']:,.2f}",
                'aggiornamento_anagrafica': f"€ {d['c2']:,.2f}",
                'adempimenti_fiscali':      f"€ {d['c3']:,.2f}",
                'certificazione_unica':     f"€ {d['c4']:,.2f}",
                'totale_imponibile':        f"€ {imp:,.2f}",
                'cassa_previdenza':         f"€ {inar:,.2f}",
                'iva':                      f"€ {iva_imp:,.2f}",
                'totale_finale':            f"€ {totale:,.2f}",
                'note':                     d["note"],
            }
            doc = DocxTemplate(path_modello)
            doc.render(ctx)
            try:
                doc.save(percorso_salvataggio)
                word_ok = True
            except PermissionError:
                raise PermissionError("FILE BLOCCATO: Chiudi il file Word aperto e riprova.")

            time.sleep(1)

            pdf_ok = converti_pdf_diretto(percorso_salvataggio, percorso_pdf)
            if not pdf_ok:
                err = "Word salvato. Conversione PDF fallita (Word bloccato o licenza?)."
        except Exception as e:
            err = str(e)

        app.after(0, lambda: end_pdf(loading, err, percorso_salvataggio, pdf_ok, word_ok))

    threading.Thread(target=task, daemon=True).start()

def end_pdf(win, e, p_w, ok, w_saved):
    win.destroy()
    if ok and w_saved:
        # M2 – incrementa il numero progressivo dopo salvataggio riuscito
        config["numero_preventivo"] = config.get("numero_preventivo", 1) + 1
        try:
            with open(path_config, 'w') as fh:
                json.dump(config, fh, indent=4)
        except:
            pass
        Messagebox.show_info("Ottimo Lavoro!",
            f"Tutto salvato in:\n{os.path.dirname(p_w)}\n\nWord e PDF pronti.")
        try: os.startfile(p_w)
        except: pass
        aggiorna_storico()
    elif e:
        Messagebox.show_warning("Attenzione", f"{e}")
        if w_saved:
            try: os.startfile(p_w)
            except: pass

# ─────────────────────────────────────────────
# Storico preventivi
# ─────────────────────────────────────────────
def aggiorna_storico(filtro=""):
    """Rilegge la cartella di salvataggio e aggiorna la lista storico."""
    try:
        listbox_storico.delete(0, END)
        cartella = config.get("ultima_cartella", "")
        if not cartella or not os.path.isdir(cartella):
            cartella = app_path
        lbl_storico_cartella.config(text=f"Cartella: {cartella}")
        files = sorted(
            [f for f in os.listdir(cartella) if f.startswith("Preventivo_") and f.endswith(".docx")],
            key=lambda x: os.path.getmtime(os.path.join(cartella, x)),
            reverse=True
        )
        # M4 – applica filtro di ricerca
        if filtro:
            files = [f for f in files if filtro.lower() in f.lower()]
        for f in files:
            mtime = os.path.getmtime(os.path.join(cartella, f))
            data  = datetime.fromtimestamp(mtime).strftime("%d/%m/%Y %H:%M")
            listbox_storico.insert(END, f"  {data}   {f}")
    except Exception as ex:
        listbox_storico.insert(END, f"Errore lettura: {ex}")

def apri_selezionato():
    selezione = listbox_storico.curselection()
    if not selezione:
        return
    riga    = listbox_storico.get(selezione[0])
    nome_f  = riga.strip().split("   ", 1)[-1]
    cartella= config.get("ultima_cartella", app_path)
    percorso= os.path.join(cartella, nome_f)
    if os.path.exists(percorso):
        os.startfile(percorso)
    else:
        Messagebox.show_error("Errore", f"File non trovato:\n{percorso}")

# M3 – eliminazione file dallo storico con conferma
def elimina_selezionato():
    selezione = listbox_storico.curselection()
    if not selezione:
        Messagebox.show_warning("Nessuna selezione", "Seleziona un file dalla lista prima di eliminare.")
        return
    riga    = listbox_storico.get(selezione[0])
    nome_f  = riga.strip().split("   ", 1)[-1]
    cartella = config.get("ultima_cartella", app_path)
    percorso = os.path.join(cartella, nome_f)
    if not os.path.exists(percorso):
        Messagebox.show_error("Errore", f"File non trovato:\n{percorso}")
        return
    risposta = Messagebox.yesno(
        "Conferma eliminazione",
        f"Vuoi eliminare definitivamente:\n\n{nome_f}\n\n(verrà rimosso anche il PDF associato, se presente)"
    )
    if risposta == "Yes":
        try:
            os.remove(percorso)
            percorso_pdf = percorso.replace(".docx", ".pdf")
            if os.path.exists(percorso_pdf):
                os.remove(percorso_pdf)
            aggiorna_storico()
        except Exception as ex:
            Messagebox.show_error("Errore", f"Impossibile eliminare:\n{ex}")

def sfoglia_cartella_storico():
    nuova = filedialog.askdirectory(title="Scegli la cartella dei preventivi",
                                   initialdir=config.get("ultima_cartella", app_path))
    if nuova:
        config["ultima_cartella"] = nuova
        try:
            with open(path_config, 'w') as fh:
                json.dump(config, fh, indent=4)
        except:
            pass
        aggiorna_storico()

# ════════════════════════════════════════
# BILANCI – Gestione date chiusura
# ════════════════════════════════════════
def carica_bilanci():
    if os.path.exists(path_bilanci):
        try:
            with open(path_bilanci, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            pass
    return []

def salva_bilanci_json(dati):
    try:
        with open(path_bilanci, 'w', encoding='utf-8') as f:
            json.dump(dati, f, indent=4, ensure_ascii=False)
    except Exception as e:
        Messagebox.show_error("Errore", f"Impossibile salvare:\n{e}")

def aggiorna_treeview_bilanci():
    for item in tree_bilanci.get_children():
        tree_bilanci.delete(item)
    bilanci = carica_bilanci()
    oggi = datetime.now().date()
    def _parse(b):
        try: return datetime.strptime(b['data_chiusura'], '%Y-%m-%d').date()
        except: return None
    for b in sorted(bilanci, key=lambda x: _parse(x) or datetime.max.date()):
        data_b = _parse(b)
        if data_b:
            g = (data_b - oggi).days
            ds = data_b.strftime('%d/%m/%Y')
            if g < 0:    gs, tag = f"⛔ Scaduto ({abs(g)}gg fa)", 'scaduto'
            elif g <= 30: gs, tag = f"🔴 {g} giorni",            'urgente'
            elif g <= 90: gs, tag = f"🟡 {g} giorni",            'presto'
            else:         gs, tag = f"✅ {g} giorni",            'ok'
        else:
            ds, gs, tag = 'N/D', '—', 'ok'
        tree_bilanci.insert('', END, values=(b.get('id',''), b.get('nome',''),
            b.get('indirizzo',''), ds, gs, b.get('note','')), tags=(tag,))

def apri_form_bilancio(esistente=None):
    dlg = ttk.Toplevel()
    dlg.title("Modifica Bilancio" if esistente else "Nuovo Bilancio")
    dlg.resizable(False, False)
    w, h = 430, 300
    dlg.geometry(f"{w}x{h}+{app.winfo_x()+(app.winfo_width()//2)-(w//2)}+{app.winfo_y()+(app.winfo_height()//2)-(h//2)}")
    dlg.grab_set()
    ttk.Label(dlg, text="Modifica Bilancio" if esistente else "Nuovo Bilancio",
              font=("Helvetica",12,"bold"), bootstyle="primary").pack(pady=(15,8))
    frm = ttk.Frame(dlg); frm.pack(fill=X, padx=20)
    frm.columnconfigure(1, weight=1)
    def mkrow(lbl, row, val=""):
        ttk.Label(frm, text=lbl, width=22, anchor=W).grid(row=row, column=0, sticky=W, pady=4)
        e = ttk.Entry(frm); e.insert(0, val); e.grid(row=row, column=1, sticky=EW, padx=5)
        return e
    e_nome = mkrow("Nome Condominio:",  0, esistente.get('nome','')       if esistente else '')
    e_ind  = mkrow("Indirizzo:",        1, esistente.get('indirizzo','')  if esistente else '')
    # data
    ttk.Label(frm, text="Chiusura Bilancio:", width=22, anchor=W).grid(row=2, column=0, sticky=W, pady=4)
    e_data = ttk.Entry(frm)
    if esistente and esistente.get('data_chiusura'):
        try: e_data.insert(0, datetime.strptime(esistente['data_chiusura'],'%Y-%m-%d').strftime('%d/%m/%Y'))
        except: pass
    e_data.grid(row=2, column=1, sticky=EW, padx=5)
    ttk.Label(frm, text="(gg/mm/aaaa)", font=("Arial",8,"italic"), bootstyle="secondary").grid(row=3, column=1, sticky=W, padx=5)
    e_note = mkrow("Note:", 4, esistente.get('note','') if esistente else '')
    def salva():
        nome = e_nome.get().strip()
        if not nome: Messagebox.show_warning("Attenzione","Inserisci il nome del condominio."); return
        try: data_iso = datetime.strptime(e_data.get().strip(),'%d/%m/%Y').strftime('%Y-%m-%d')
        except: Messagebox.show_error("Errore","Data non valida (usa gg/mm/aaaa)."); return
        bilanci = carica_bilanci()
        if esistente:
            for b in bilanci:
                if b['id'] == esistente['id']:
                    b.update({'nome':nome,'indirizzo':e_ind.get().strip(),'data_chiusura':data_iso,'note':e_note.get().strip()})
                    break
        else:
            bilanci.append({'id': max((b['id'] for b in bilanci), default=0)+1,
                            'nome':nome,'indirizzo':e_ind.get().strip(),
                            'data_chiusura':data_iso,'note':e_note.get().strip()})
        salva_bilanci_json(bilanci)
        dlg.destroy()
        aggiorna_treeview_bilanci()
    fr_btn = ttk.Frame(dlg); fr_btn.pack(fill=X, padx=20, pady=12)
    ttk.Button(fr_btn, text="💾 Salva",  command=salva,       bootstyle="success",          width=15).pack(side=LEFT,  expand=True, fill=X, padx=5)
    ttk.Button(fr_btn, text="Annulla",   command=dlg.destroy, bootstyle="secondary-outline",width=12).pack(side=RIGHT, padx=5)

def modifica_bilancio_selezionato():
    sel = tree_bilanci.selection()
    if not sel: Messagebox.show_warning("Nessuna selezione","Seleziona un bilancio."); return
    bid = tree_bilanci.item(sel[0])['values'][0]
    b = next((x for x in carica_bilanci() if x['id']==bid), None)
    if b: apri_form_bilancio(b)

def elimina_bilancio_selezionato():
    sel = tree_bilanci.selection()
    if not sel: Messagebox.show_warning("Nessuna selezione","Seleziona un bilancio."); return
    vals = tree_bilanci.item(sel[0])['values']
    if Messagebox.yesno("Conferma", f"Eliminare il bilancio di:\n{vals[1]}?") == "Yes":
        bilanci = [b for b in carica_bilanci() if b['id'] != vals[0]]
        salva_bilanci_json(bilanci)
        aggiorna_treeview_bilanci()

def stampa_bilanci():
    ultima = config.get("ultima_cartella","")
    if not ultima or not os.path.isdir(ultima): ultima = app_path
    percorso = filedialog.asksaveasfilename(
        title="Salva Report Bilanci", initialdir=ultima,
        initialfile=f"ScadenzeBilanci_{datetime.now().strftime('%Y%m%d')}.docx",
        defaultextension=".docx", filetypes=[("Word Document","*.docx")])
    if not percorso: return
    loading = mostra_caricamento("Generazione report bilanci...")
    def task():
        try: pythoncom.CoInitialize()
        except: pass
        err = None; pdf_ok = False
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            bilanci = carica_bilanci()
            oggi = datetime.now().date()
            def _parse(b):
                try: return datetime.strptime(b['data_chiusura'],'%Y-%m-%d').date()
                except: return None
            bilanci_s = sorted(bilanci, key=lambda x: _parse(x) or datetime.max.date())
            doc = Document()
            sec = doc.sections[0]
            sec.left_margin=sec.right_margin=Cm(2); sec.top_margin=sec.bottom_margin=Cm(2)
            # Titolo
            for txt, sz, bold, italic in [
                (config.get('nome_studio','STUDIO SERRA'), 18, True,  False),
                ('SCADENZE CHIUSURA BILANCI CONDOMINIALI', 14, True,  False),
                (f"Aggiornato al: {oggi.strftime('%d/%m/%Y')}", 10, False, True),
            ]:
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt); r.font.size=Pt(sz); r.bold=bold; r.italic=italic
                if bold and sz==18: r.font.color.rgb=RGBColor(0x1A,0x5C,0x96)
            doc.add_paragraph()
            # Sezione urgenti
            urgenti = [b for b in bilanci_s if _parse(b) and ((_parse(b)-oggi).days) < 90]
            if urgenti:
                p=doc.add_paragraph(); r=p.add_run('⚠  SCADENZE IMMINENTI / SCADUTE')
                r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor(0xCC,0,0)
                for b in urgenti:
                    db=_parse(b); g=(db-oggi).days
                    stato=f'SCADUTO da {abs(g)}gg' if g<0 else f'tra {g} giorni'
                    pi=doc.add_paragraph(style='List Bullet')
                    ri=pi.add_run(f"{b['nome']}  —  {db.strftime('%d/%m/%Y')}  ({stato})")
                    ri.font.color.rgb=RGBColor(0xCC,0,0) if g<0 else RGBColor(0xE6,0x7E,0x22)
                doc.add_paragraph()
            # Tabella
            p=doc.add_paragraph(); r=p.add_run('ELENCO COMPLETO'); r.bold=True; r.font.size=Pt(12)
            table=doc.add_table(rows=1,cols=5); table.style='Table Grid'
            hdrs=['Condominio','Indirizzo','Chiusura','Giorni Mancanti','Note']
            ws=[Cm(5.5),Cm(5),Cm(3),Cm(3),Cm(2.5)]
            def shd_cell(cell, fill):
                tc=cell._tc; tcPr=tc.get_or_add_tcPr()
                s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear')
                s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),fill); tcPr.append(s)
            for i,(h,w) in enumerate(zip(hdrs,ws)):
                c=table.rows[0].cells[i]; c.width=w
                c.paragraphs[0].clear()
                r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(9)
                r.font.color.rgb=RGBColor(255,255,255); shd_cell(c,'1A5C96')
                c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
            for b in bilanci_s:
                db=_parse(b)
                if db:
                    g=(db-oggi).days; ds=db.strftime('%d/%m/%Y')
                    if g<0:    gs,col,fill=f'SCADUTO ({abs(g)}gg)',RGBColor(0xCC,0,0),'FFE6E6'
                    elif g<=30: gs,col,fill=f'{g} giorni',RGBColor(0xCC,0x44,0),'FFF0CC'
                    elif g<=90: gs,col,fill=f'{g} giorni',RGBColor(0x88,0x66,0),'FFFAEE'
                    else:       gs,col,fill=f'{g} giorni',RGBColor(0x1A,0x7A,0x1A),'FFFFFF'
                else: ds,gs,col,fill='N/D','—',RGBColor(0,0,0),'FFFFFF'
                row=table.add_row()
                for i,val in enumerate([b.get('nome',''),b.get('indirizzo',''),ds,gs,b.get('note','')]):
                    c=row.cells[i]; c.paragraphs[0].clear()
                    r=c.paragraphs[0].add_run(val); r.font.size=Pt(9); r.font.color.rgb=col
                    if fill!='FFFFFF': shd_cell(c,fill)
            doc.add_paragraph()
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(f'Totale condomini in gestione: {len(bilanci)}')
            r.font.size=Pt(10); r.italic=True
            doc.save(percorso)
            pdf_out=percorso.replace('.docx','.pdf')
            time.sleep(0.5)
            pdf_ok=converti_pdf_diretto(percorso, pdf_out)
            if not pdf_ok: err='Word salvato. Conversione PDF fallita.'
        except Exception as e:
            err=str(e)
        app.after(0, lambda: _fine_bilanci(loading, err, percorso, pdf_ok))
    threading.Thread(target=task, daemon=True).start()

def _fine_bilanci(win, err, p_w, pdf_ok):
    win.destroy()
    if err: Messagebox.show_warning("Attenzione", err)
    else:
        Messagebox.show_info("Fatto!",f"Report salvato in:\n{os.path.dirname(p_w)}")
        try: os.startfile(p_w)
        except: pass

# ════════════════════════════════════════
# --- 6. INTERFACCIA ---
# ════════════════════════════════════════
try:
    app = ttk.Window(themename="lumen")
    app.title("Gestione Preventivi - Studio Serra v4.1")
    app.geometry("700x820")

    try:
        app.iconbitmap(path_icona)
    except:
        pass

    # ---- HEADER ----
    fr_head = ttk.Frame(app); fr_head.pack(fill=X, pady=(10, 0))
    try:
        if os.path.exists(path_logo):
            img = ttk.PhotoImage(file=path_logo)
            ttk.Label(fr_head, image=img).pack(side=LEFT, padx=10)
        else:
            ttk.Label(fr_head, text=config["nome_studio"],
                      font=("Helvetica", 20, "bold"), bootstyle="primary").pack(side=LEFT, padx=10)
    except:
        pass
    ttk.Button(fr_head, text="🗑️  Nuovo", command=nuovo_preventivo,
               bootstyle="secondary-outline").pack(side=RIGHT, padx=10, pady=5)

    # ---- NOTEBOOK ----
    nb = ttk.Notebook(app); nb.pack(fill=BOTH, expand=True, padx=10, pady=5)
    t1 = ttk.Frame(nb); nb.add(t1, text="  PREVENTIVO  ")
    t2 = ttk.Frame(nb); nb.add(t2, text="  STORICO  ")
    t3 = ttk.Frame(nb); nb.add(t3, text="  CONFIGURAZIONE  ")
    t4 = ttk.Frame(nb); nb.add(t4, text="  📅 BILANCI  ")

    # ════════════════════════════════════════
    # TAB 1 – PREVENTIVO
    # ════════════════════════════════════════
    cv = tk.Canvas(t1); sc = ttk.Scrollbar(t1, command=cv.yview)
    pc = ttk.Frame(cv)
    pc.bind("<Configure>", lambda e: cv.configure(scrollregion=cv.bbox("all")) if e.widget == pc else None)
    cv.create_window((0, 0), window=pc, anchor="nw")
    cv.configure(yscrollcommand=sc.set)
    cv.pack(side=LEFT, fill=BOTH, expand=True); sc.pack(side=RIGHT, fill=Y)

    def _on_mousewheel(event):
        if cv.winfo_height() < pc.winfo_height():
            cv.yview_scroll(int(-1 * (event.delta / 120)), "units")
    cv.bind("<MouseWheel>", _on_mousewheel)

    box = ttk.Frame(pc); box.pack(fill=X, padx=10, pady=10)

    # 1. Dati Condominio
    lf1 = ttk.LabelFrame(box, text="  Dati Condominio  "); lf1.pack(fill=X, pady=5)
    in1 = ttk.Frame(lf1); in1.pack(fill=X, padx=10, pady=10)
    ttk.Label(in1, text="Nome Condominio:").pack(anchor=W)
    entry_nome = ttk.Entry(in1); entry_nome.pack(fill=X)
    ttk.Label(in1, text="Indirizzo:").pack(anchor=W, pady=(5, 0))
    entry_indirizzo = ttk.Entry(in1); entry_indirizzo.pack(fill=X)
    ttk.Label(in1, text="Città:").pack(anchor=W, pady=(5, 0))
    entry_citta = ttk.Entry(in1); entry_citta.insert(0, "Bergamo"); entry_citta.pack(fill=X)

    # 2. Dati Tecnici
    lf2 = ttk.LabelFrame(box, text="  Dati Tecnici  "); lf2.pack(fill=X, pady=5)
    in2 = ttk.Frame(lf2); in2.pack(fill=X, padx=10, pady=10)
    gr = ttk.Frame(in2); gr.pack(fill=X)
    ttk.Label(gr, text="Spesa totale (€):").grid(row=0, column=0, sticky=W)
    entry_spesa = ttk.Entry(gr, width=10); entry_spesa.insert(0, "0"); entry_spesa.grid(row=0, column=1, padx=5)
    ttk.Label(gr, text="Ascensori:").grid(row=0, column=2, sticky=W)
    entry_ascensori = ttk.Entry(gr, width=5); entry_ascensori.insert(0, "0"); entry_ascensori.grid(row=0, column=3, padx=5)
    ttk.Label(gr, text="Cancelli:").grid(row=0, column=4, sticky=W)
    entry_cancelli = ttk.Entry(gr, width=5); entry_cancelli.insert(0, "0"); entry_cancelli.grid(row=0, column=5, padx=5)

    fr_c = ttk.Frame(in2); fr_c.pack(fill=X, pady=5)
    var_risc = ttk.IntVar(); ttk.Checkbutton(fr_c, text="Riscaldamento", variable=var_risc, command=aggiorna_riepilogo).pack(side=LEFT, padx=5)
    var_port = ttk.IntVar(); ttk.Checkbutton(fr_c, text="Portiere",       variable=var_port, command=aggiorna_riepilogo).pack(side=LEFT, padx=5)
    var_pisc = ttk.IntVar(); ttk.Checkbutton(fr_c, text="Piscina",        variable=var_pisc, command=aggiorna_riepilogo).pack(side=LEFT, padx=5)

    # M5 – mostra formula trasferte accanto al campo km
    fr_d = ttk.Frame(in2); fr_d.pack(fill=X)
    ttk.Label(fr_d, text="Km (A/R x4):").pack(side=LEFT)
    entry_km = ttk.Entry(fr_d, width=8); entry_km.insert(0, "0"); entry_km.pack(side=LEFT, padx=5)
    ttk.Button(fr_d, text="📍 Calcola distanza", command=avvia_mappe, bootstyle="info-outline").pack(side=LEFT)
    ttk.Label(fr_d, text=f"  ⟹  km × 4 × 2 × €{config['costo_km']}/km",
              font=("Arial", 8, "italic"), bootstyle="secondary").pack(side=LEFT, padx=5)

    # 3. Economico
    lf3 = ttk.LabelFrame(box, text="  Economico  "); lf3.pack(fill=X, pady=5)
    in3 = ttk.Frame(lf3); in3.pack(fill=X, padx=10, pady=10)
    r1 = ttk.Frame(in3); r1.pack(fill=X)
    ttk.Label(r1, text="Unità abitative:").pack(side=LEFT)
    entry_unita = ttk.Entry(r1, width=8); entry_unita.pack(side=LEFT, padx=5)
    entry_unita.bind("<KeyRelease>", calcola_tariffa_automatica)
    ttk.Label(r1, text="Box auto (5 box = 1 unità):").pack(side=LEFT)
    entry_box = ttk.Entry(r1, width=8); entry_box.insert(0, "0"); entry_box.pack(side=LEFT, padx=5)

    r2 = ttk.Frame(in3); r2.pack(fill=X, pady=5)
    ttk.Label(r2, text="Prezzo/unità (€):").pack(side=LEFT)
    entry_prezzo_unita = ttk.Entry(r2, width=8); entry_prezzo_unita.insert(0, "0"); entry_prezzo_unita.pack(side=LEFT, padx=5)
    lbl_tariffa = ttk.Label(r2, text="", font=("Arial", 9, "italic")); lbl_tariffa.pack(side=LEFT)

    r3 = ttk.Frame(in3); r3.pack(fill=X)
    ttk.Label(r3, text="Sconto (%):").pack(side=LEFT)
    entry_sconto = ttk.Entry(r3, width=8); entry_sconto.insert(0, "0"); entry_sconto.pack(side=LEFT, padx=5)

    # 4. Spese Fisse
    lf4 = ttk.LabelFrame(box, text="  Spese Fisse  "); lf4.pack(fill=X, pady=5)
    in4 = ttk.Frame(lf4); in4.pack(fill=X, padx=10, pady=10)
    def mkf(parent, testo, chiave, riga, col):
        ttk.Label(parent, text=testo).grid(row=riga, column=col, sticky=W)
        e = ttk.Entry(parent, width=9); e.insert(0, str(config[chiave])); e.grid(row=riga, column=col + 1, padx=5, pady=3)
        return e
    entry_canc = mkf(in4, "Cancelleria (€):",    "costo_cancelleria", 0, 0)
    entry_adem = mkf(in4, "Adempimenti fiscali:", "costo_fiscali",     0, 2)
    entry_anag = mkf(in4, "Anagrafica (€):",      "costo_anagrafica",  1, 0)
    entry_cu   = mkf(in4, "Certif. Unica (€):",   "costo_cu",          1, 2)

    # 5. Note
    lf5 = ttk.LabelFrame(box, text="  Note / Condizioni Particolari  "); lf5.pack(fill=X, pady=5)
    in5 = ttk.Frame(lf5); in5.pack(fill=X, padx=10, pady=5)
    entry_note = tk.Text(in5, height=3, font=("Arial", 10), relief="flat",
                         bg="#f8f9fa", bd=1, highlightthickness=1, highlightbackground="#ced4da")
    entry_note.pack(fill=X)

    # ─────────────────────────────────────────────
    # RIEPILOGO IMPORTI in tempo reale
    # ─────────────────────────────────────────────
    lf_rie = ttk.LabelFrame(box, text="  📊 Riepilogo Importi  "); lf_rie.pack(fill=X, pady=5)
    fr_rie = ttk.Frame(lf_rie); fr_rie.pack(fill=X, padx=10, pady=8)

    def _riga_rie(parent, testo, riga, stile="default"):
        ttk.Label(parent, text=testo, font=("Arial", 10)).grid(row=riga, column=0, sticky=W, padx=5, pady=2)
        lbl = ttk.Label(parent, text="€ 0,00", font=("Arial", 10, "bold"), bootstyle=stile, width=14, anchor=E)
        lbl.grid(row=riga, column=1, sticky=E, padx=5)
        return lbl

    lbl_riepilogo_imponibile = _riga_rie(fr_rie, "Totale Imponibile:",   0, "info")
    lbl_riepilogo_cassa      = _riga_rie(fr_rie, "Cassa Previdenza 4%:", 1, "warning")
    lbl_riepilogo_iva        = _riga_rie(fr_rie, "IVA 22%:",              2, "secondary")
    lbl_riepilogo_totale     = _riga_rie(fr_rie, "TOTALE FINALE:",        3, "success")
    fr_rie.columnconfigure(0, weight=1)

    # Pulsante principale
    ttk.Button(box, text="💾  SALVA PREVENTIVO (WORD + PDF)",
               command=avvia_pdf, bootstyle="success", width=36).pack(pady=15)

    # Binding riepilogo su tutti i campi numerici
    for _e in (entry_spesa, entry_ascensori, entry_cancelli,
               entry_box, entry_km, entry_prezzo_unita, entry_sconto,
               entry_canc, entry_anag, entry_adem, entry_cu):
        _e.bind("<KeyRelease>", aggiorna_riepilogo)

    def _setup_campo_numerico(entry_widget, default_val="0"):
        def _on_focus_in(event):
            v = entry_widget.get().strip()
            if v in ("0", "0.0", "0.00"):
                entry_widget.delete(0, tk.END)
            else:
                entry_widget.select_range(0, tk.END)
                entry_widget.icursor(tk.END)

        def _on_focus_out(event):
            v = entry_widget.get().strip()
            if not v:
                entry_widget.insert(0, default_val)
                aggiorna_riepilogo()

        entry_widget.bind("<FocusIn>", _on_focus_in)
        entry_widget.bind("<FocusOut>", _on_focus_out)

    def _setup_campo_testo_default(entry_widget, default_val="Bergamo"):
        def _on_focus_in(event):
            v = entry_widget.get().strip()
            if v == default_val:
                entry_widget.select_range(0, tk.END)
                entry_widget.icursor(tk.END)

        entry_widget.bind("<FocusIn>", _on_focus_in)

    # Gestione automatica azzeramento/selezione al focus per digitazione fluida
    for _e in (entry_spesa, entry_ascensori, entry_cancelli, entry_km,
               entry_box, entry_prezzo_unita, entry_sconto,
               entry_canc, entry_anag, entry_adem, entry_cu):
        _setup_campo_numerico(_e, "0")

    _setup_campo_testo_default(entry_citta, "Bergamo")

    # ════════════════════════════════════════
    # TAB 2 – STORICO
    # ════════════════════════════════════════
    fr_sto = ttk.Frame(t2); fr_sto.pack(fill=BOTH, expand=True, padx=10, pady=10)

    fr_sto_top = ttk.Frame(fr_sto); fr_sto_top.pack(fill=X, pady=5)
    lbl_storico_cartella = ttk.Label(fr_sto_top, text="Cartella: —", font=("Arial", 9), bootstyle="secondary")
    lbl_storico_cartella.pack(side=LEFT, fill=X, expand=True)
    ttk.Button(fr_sto_top, text="📁 Cambia cartella", command=sfoglia_cartella_storico,
               bootstyle="info-outline").pack(side=RIGHT)
    ttk.Button(fr_sto_top, text="🔄 Aggiorna", command=aggiorna_storico,
               bootstyle="secondary-outline").pack(side=RIGHT, padx=5)

    # M4 – campo di ricerca/filtro
    fr_sto_search = ttk.Frame(fr_sto); fr_sto_search.pack(fill=X, pady=(0, 5))
    ttk.Label(fr_sto_search, text="🔍 Cerca:", font=("Arial", 10)).pack(side=LEFT)
    entry_storico_filtro = ttk.Entry(fr_sto_search, font=("Arial", 10))
    entry_storico_filtro.pack(side=LEFT, fill=X, expand=True, padx=5)

    def _filtra_storico(*args):
        aggiorna_storico(filtro=entry_storico_filtro.get())
    entry_storico_filtro.bind("<KeyRelease>", _filtra_storico)

    fr_list = ttk.Frame(fr_sto); fr_list.pack(fill=BOTH, expand=True)
    scrollbar_sto = ttk.Scrollbar(fr_list)
    scrollbar_sto.pack(side=RIGHT, fill=Y)
    listbox_storico = tk.Listbox(fr_list, yscrollcommand=scrollbar_sto.set,
                                 font=("Courier New", 10), selectmode=SINGLE,
                                 relief="flat", bd=0, activestyle="dotbox", height=20)
    listbox_storico.pack(fill=BOTH, expand=True)
    scrollbar_sto.config(command=listbox_storico.yview)

    # Pulsanti azioni storico
    fr_sto_btn = ttk.Frame(fr_sto); fr_sto_btn.pack(fill=X, pady=5)
    ttk.Button(fr_sto_btn, text="📂 Apri file selezionato", command=apri_selezionato,
               bootstyle="primary").pack(side=LEFT, fill=X, expand=True, padx=(0, 5))
    # M3 – pulsante elimina
    ttk.Button(fr_sto_btn, text="🗑️  Elimina", command=elimina_selezionato,
               bootstyle="danger-outline").pack(side=RIGHT)

    # ════════════════════════════════════════
    # TAB 3 – CONFIGURAZIONE
    # ════════════════════════════════════════
    cfbox = ttk.Frame(t3); cfbox.pack(fill=BOTH, expand=True, padx=20, pady=20)
    def mkc(parent, testo, chiave):
        f = ttk.Frame(parent); f.pack(fill=X, pady=3)
        ttk.Label(f, text=testo, width=22).pack(side=LEFT)
        e = ttk.Entry(f); e.insert(0, str(config[chiave])); e.pack(side=RIGHT, fill=X, expand=True)
        return e
    ent_cfg_nome = mkc(cfbox, "Nome Studio:",        "nome_studio")
    ent_cfg_addr = mkc(cfbox, "Indirizzo Studio:",   "indirizzo_studio")
    ttk.Separator(cfbox).pack(fill=X, pady=8)
    ttk.Label(cfbox, text="Tariffe per fascia (€/unità)", font=("Arial", 10, "bold")).pack(anchor=W)
    ent_cfg_p1 = mkc(cfbox, "Fascia 1 (< 10 unità):",  "prezzo_fascia_1")
    ent_cfg_p2 = mkc(cfbox, "Fascia 2 (10-20 unità):", "prezzo_fascia_2")
    ent_cfg_p3 = mkc(cfbox, "Fascia 3 (21-30 unità):", "prezzo_fascia_3")
    ent_cfg_p4 = mkc(cfbox, "Fascia 4 (> 30 unità):",  "prezzo_fascia_4")
    ttk.Separator(cfbox).pack(fill=X, pady=8)
    ttk.Label(cfbox, text="Costi fissi default (€)", font=("Arial", 10, "bold")).pack(anchor=W)
    ent_cfg_canc = mkc(cfbox, "Cancelleria:",         "costo_cancelleria")
    ent_cfg_anag = mkc(cfbox, "Anagrafica:",          "costo_anagrafica")
    ent_cfg_fisc = mkc(cfbox, "Adempimenti fiscali:", "costo_fiscali")
    ent_cfg_cu   = mkc(cfbox, "Certif. Unica:",       "costo_cu")
    ttk.Separator(cfbox).pack(fill=X, pady=8)
    # M5 – costo_km nella configurazione
    ttk.Label(cfbox, text="Trasferte", font=("Arial", 10, "bold")).pack(anchor=W)
    ent_cfg_km = mkc(cfbox, "Costo km (€/km):",      "costo_km")
    ttk.Label(cfbox, text="Formula: km × 4 uscite × 2 (A/R) × costo km",
              font=("Arial", 8, "italic"), bootstyle="secondary").pack(anchor=W, pady=(0, 4))
    ttk.Button(cfbox, text="💾  SALVA CONFIGURAZIONE",
               command=salva_configurazione, bootstyle="primary").pack(pady=15, fill=X)

    # ════════════════════════════════════════
    # TAB 4 – BILANCI
    # ════════════════════════════════════════
    fr_bil = ttk.Frame(t4); fr_bil.pack(fill=BOTH, expand=True, padx=10, pady=10)

    fr_bil_top = ttk.Frame(fr_bil); fr_bil_top.pack(fill=X, pady=5)
    ttk.Label(fr_bil_top, text="Scadenze Chiusura Bilanci Condominiali",
              font=("Helvetica",12,"bold"), bootstyle="primary").pack(side=LEFT)
    ttk.Button(fr_bil_top, text="📄 Stampa PDF", command=stampa_bilanci,
               bootstyle="success").pack(side=RIGHT, padx=5)
    ttk.Button(fr_bil_top, text="➕ Aggiungi", command=lambda: apri_form_bilancio(),
               bootstyle="primary-outline").pack(side=RIGHT, padx=5)

    # Legenda
    fr_leg = ttk.Frame(fr_bil); fr_leg.pack(fill=X, pady=2)
    for txt in ["⛔ Scaduto", "  🔴 Entro 30gg", "  🟡 Entro 90gg", "  ✅ Oltre 90gg"]:
        ttk.Label(fr_leg, text=txt, font=("Arial",9), bootstyle="secondary").pack(side=LEFT)

    # Treeview
    fr_tree = ttk.Frame(fr_bil); fr_tree.pack(fill=BOTH, expand=True, pady=5)
    sb_bil = ttk.Scrollbar(fr_tree); sb_bil.pack(side=RIGHT, fill=Y)
    tree_bilanci = ttk.Treeview(fr_tree, yscrollcommand=sb_bil.set,
                                columns=('id','nome','indirizzo','data','giorni','note'),
                                show='headings', selectmode='browse', height=18)
    tree_bilanci.pack(fill=BOTH, expand=True)
    sb_bil.config(command=tree_bilanci.yview)

    for col, hdr, w, anchor in [
        ('id','',0,'center'),('nome','Condominio',180,'w'),
        ('indirizzo','Indirizzo',150,'w'),('data','Chiusura Bilancio',110,'center'),
        ('giorni','Giorni Mancanti',120,'center'),('note','Note',140,'w')
    ]:
        tree_bilanci.heading(col, text=hdr)
        tree_bilanci.column(col, width=w, anchor=anchor, stretch=(col!='id'))

    tree_bilanci.tag_configure('scaduto', foreground='#cc0000', background='#ffe6e6')
    tree_bilanci.tag_configure('urgente', foreground='#cc4400', background='#fff0cc')
    tree_bilanci.tag_configure('presto',  foreground='#886600', background='#fffaee')
    tree_bilanci.tag_configure('ok',      foreground='#1a7a1a')

    # Pulsanti azioni
    fr_bil_btn = ttk.Frame(fr_bil); fr_bil_btn.pack(fill=X, pady=5)
    ttk.Button(fr_bil_btn, text="✏️ Modifica Selezionato",
               command=modifica_bilancio_selezionato,
               bootstyle="info-outline").pack(side=LEFT, padx=5)
    ttk.Button(fr_bil_btn, text="🗑️ Elimina Selezionato",
               command=elimina_bilancio_selezionato,
               bootstyle="danger-outline").pack(side=LEFT)

    # Avvio
    aggiorna_riepilogo()
    aggiorna_storico()
    aggiorna_treeview_bilanci()

    app.mainloop()

except Exception as e:
    show_error_msg("Crash", str(e))
